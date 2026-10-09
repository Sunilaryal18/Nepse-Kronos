"""Collect monthly macro data from Nepal Rastra Bank (NRB) into data/nepse/meta/macro_monthly.csv.

Sources (newest release of each; every release carries the full history since Aug 2016):
  * "Current Macroeconomic and Financial Situation - Tables" xlsx
      - sheet "A22.Interest Rates- monthly": interbank, 91-day T-bill, weighted average
        deposit and lending rates, base rate (commercial banks)
      - sheet "A7..BOP- Annex-Monthly": Workers' Remittances, Rs million, cumulative within the
        Nepali fiscal year (which starts with Shrawan, mid-July) -> de-cumulated to monthly
  * "Banking and Financial Statistics" (monthly statistics) xlsx
      - sheet "C18", Table A: Total Credit/Total Deposit (%)

Month convention
  NRB labels each Nepali (Bikram Sambat, BS) month by the AD month in which it ENDS, e.g. the
  Shrawan 2083 figures are published as "Mid-August 2026" / fiscal year 2026/27, "August".
  We store that as ``month`` = the 16th of the labelled AD month (2026-08-16), a fixed stand-in
  for the BS month-end (which falls between the 14th and 17th). In a fiscal-year row "YYYY/YY",
  August..December belong to the first AD year and January..July to the second.
  BS month -> AD reference month: Shrawan->Aug, Bhadra->Sep, Ashwin->Oct, Kartik->Nov,
  Mangsir->Dec, Poush->Jan, Magh->Feb, Falgun->Mar, Chaitra->Apr, Baisakh->May, Jestha->Jun,
  Ashadh->Jul; AD year = BS year - 57 for Baisakh..Mangsir and BS year - 56 for Poush..Chaitra.

No look-ahead: ``available_from`` = month + 40 days (NRB releases lag about five weeks).

Usage:
    python -m nepse_kronos.nrb                 # download newest releases (cached) and write CSV
    python -m nepse_kronos.nrb --offline       # reuse cached listings/workbooks only
"""
import argparse
import html
import re
import time
from pathlib import Path
from urllib.parse import urljoin

import numpy as np
import pandas as pd

BASE = "https://www.nrb.org.np"
CMES_FEED = BASE + "/category/current-macroeconomic-situation/feed/?paged={page}"
BFS_PAGE = BASE + "/category/monthly-statistics/"
USER_AGENT = "Kronos-NEPSE-research/1.0 (personal research; polite crawler, 1 req/s)"
DEFAULT_CACHE = Path("data/nepse/source/nrb")
DEFAULT_OUT = Path("data/nepse/meta/macro_monthly.csv")

INTEREST_SHEET = "A22"
BOP_SHEET = "A7."
CD_SHEET = "C18"
AVAILABLE_LAG_DAYS = 40
REFERENCE_DAY = 16
COLUMNS = ["month", "interbank", "tbill91", "deposit_rate", "lending_rate", "base_rate",
           "cd_ratio", "remittance", "available_from"]

AD_MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august",
             "september", "october", "november", "december"]
# BS month (1 = Baisakh) -> AD month in which it ends (NRB's "mid-month" label).
BS_MONTHS = {
    "baisakh": 1, "baishakh": 1, "jestha": 2, "jeth": 2, "ashadh": 3, "asar": 3, "asadh": 3, "ashar": 3,
    "shrawan": 4, "saun": 4, "sawan": 4, "bhadra": 5, "bhadau": 5, "ashwin": 6, "asoj": 6, "ashoj": 6,
    "kartik": 7, "mangsir": 8, "mangshir": 8, "marga": 8, "poush": 9, "push": 9, "paush": 9,
    "magh": 10, "falgun": 11, "phalgun": 11, "fagun": 11, "chaitra": 12, "chait": 12,
}
BS_TO_AD_MONTH = {1: 5, 2: 6, 3: 7, 4: 8, 5: 9, 6: 10, 7: 11, 8: 12, 9: 1, 10: 2, 11: 3, 12: 4}


# ---------------------------------------------------------------- month helpers
def reference_date(year, month):
    return pd.Timestamp(year=int(year), month=int(month), day=REFERENCE_DAY)


def fiscal_month_date(fiscal_year, month_name):
    """('2016/17', 'August') -> 2016-08-16; ('2016/17', 'January') -> 2017-01-16."""
    start = int(re.match(r"\s*(\d{4})", str(fiscal_year)).group(1))
    m = AD_MONTHS.index(str(month_name).strip().lower()) + 1
    return reference_date(start if m >= 8 else start + 1, m)


def bs_month_to_reference(bs_year, bs_month):
    """BS year + month (name or 1..12, 1 = Baisakh) -> AD reference date, e.g. (2083, 'Shrawan') -> 2026-08-16."""
    if not isinstance(bs_month, (int, np.integer)):
        bs_month = BS_MONTHS[str(bs_month).strip().lower()]
    ad_month = BS_TO_AD_MONTH[int(bs_month)]
    offset = 56 if bs_month in (9, 10, 11, 12) else 57
    return reference_date(int(bs_year) - offset, ad_month)


def fiscal_year_start(month):
    """AD year in which the Nepali fiscal year containing reference date ``month`` starts."""
    return month.year if month.month >= 8 else month.year - 1


# ---------------------------------------------------------------- parsing
def _to_number(value):
    if value is None:
        return np.nan
    if isinstance(value, (int, float, np.number)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return np.nan  # "-", "", "n.a." etc.


def _find_sheet(workbook, prefix):
    for name in workbook.sheetnames:
        if name.strip().lower().startswith(prefix.lower()):
            return workbook[name]
    raise KeyError(f"no sheet starting with {prefix!r} in {workbook.sheetnames}")


def _norm(text):
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def _cell(row, col):
    return row[col] if col < len(row) else None


def parse_monthly_table(sheet, wanted):
    """Parse an NRB 'Fiscal Year | Mid-Month | ...' table.

    ``wanted`` maps output column -> substring of the header text (case/whitespace-insensitive).
    Returns a frame indexed by reference date. Stops at the first row after the data whose
    month cell is not an AD month name (footnotes, a second table, ...)."""
    rows = list(sheet.iter_rows(values_only=True))
    header_idx = None
    for i, row in enumerate(rows):
        header = [_norm(c) for c in row]
        if "fiscal year" in header and any(c.startswith("mid-month") for c in header):
            header_idx = i
            fy_col = header.index("fiscal year")
            month_col = next(j for j, c in enumerate(header) if c.startswith("mid-month"))
            break
    if header_idx is None:
        raise ValueError(f"sheet {sheet.title!r}: no 'Fiscal Year' / 'Mid-Month' header row")
    cols = {}
    for out, needle in wanted.items():
        hits = [j for j, c in enumerate(header) if _norm(needle) in c]
        if not hits:
            raise ValueError(f"sheet {sheet.title!r}: no column matching {needle!r}")
        cols[out] = hits[0]

    records, fiscal_year, started = [], None, False
    for row in rows[header_idx + 1:]:
        fy = _cell(row, fy_col)
        name = _norm(_cell(row, month_col))
        if fy is not None and re.match(r"\s*\d{4}/\d{2}", str(fy)):
            fiscal_year = str(fy).strip()
        if name not in AD_MONTHS:
            if started:
                break
            continue
        if fiscal_year is None:
            continue
        started = True
        rec = {"month": fiscal_month_date(fiscal_year, name)}
        rec.update({out: _to_number(_cell(row, j)) for out, j in cols.items()})
        records.append(rec)
    df = pd.DataFrame(records, columns=["month", *wanted])
    return df.drop_duplicates("month", keep="last").set_index("month").sort_index()


def parse_interest_rates(workbook):
    return parse_monthly_table(_find_sheet(workbook, INTEREST_SHEET), {
        "interbank": "interbank",
        "tbill91": "91 days",
        "deposit_rate": "deposit rate",
        "lending_rate": "lending rate",
        "base_rate": "base rate",
    })


def parse_cd_ratio(workbook):
    return parse_monthly_table(_find_sheet(workbook, CD_SHEET), {"cd_ratio": "total credit/total deposit"})


def parse_remittance_cumulative(workbook):
    """Workers' Remittances 'Amount' (Rs million, cumulative within the fiscal year).

    The header cell spans 'Amount' and '% Change'; 'Amount' is the first sub-column."""
    return parse_monthly_table(_find_sheet(workbook, BOP_SHEET), {"remittance_cum": "remittance"})


def decumulate_fiscal(cumulative):
    """Turn a within-fiscal-year cumulative series (index = reference dates) into monthly flows.

    The first month of each fiscal year (August reference = Shrawan) is taken as-is; a month
    whose predecessor in the same fiscal year is missing becomes NaN."""
    s = cumulative.sort_index()
    out = pd.Series(np.nan, index=s.index, name=s.name)
    for month, value in s.items():
        if month.month == 8:
            out[month] = value
            continue
        prev = (month - pd.DateOffset(months=1)).replace(day=REFERENCE_DAY)
        if prev in s.index:
            out[month] = value - s[prev]
    return out


def build_macro(tables_wb, bfs_wb):
    rates = parse_interest_rates(tables_wb)
    remit = decumulate_fiscal(parse_remittance_cumulative(tables_wb)["remittance_cum"]).rename("remittance")
    cd = parse_cd_ratio(bfs_wb)
    df = rates.join(cd, how="outer").join(remit, how="outer").sort_index().reset_index()
    df["available_from"] = df["month"] + pd.Timedelta(days=AVAILABLE_LAG_DAYS)
    df = df.dropna(subset=[c for c in COLUMNS if c not in ("month", "available_from")], how="all")
    return df[COLUMNS].reset_index(drop=True)


# ---------------------------------------------------------------- download
class PoliteSession:
    def __init__(self, delay=1.0):
        import requests
        self.session = requests.Session()
        self.session.headers["User-Agent"] = USER_AGENT
        self.delay = delay
        self._last = 0.0

    def resolve(self, url):
        """Final URL after redirects (HEAD request, so the body is not downloaded twice)."""
        return self.get(url, method="head", allow_redirects=True).url

    def get(self, url, retries=3, method="get", **kw):
        for attempt in range(retries):
            wait = self.delay - (time.monotonic() - self._last)
            if wait > 0:
                time.sleep(wait)
            self._last = time.monotonic()
            try:
                resp = self.session.request(method, url, timeout=60, **kw)
                if resp.status_code < 500:
                    resp.raise_for_status()
                    return resp
            except Exception:
                if attempt == retries - 1:
                    raise
            time.sleep(2 ** (attempt + 1))
        resp.raise_for_status()
        return resp


def _cached_text(session, url, path, offline):
    if offline or session is None:
        return path.read_text(encoding="utf-8")
    text = session.get(url).text
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return text


def _cache_file(url, cache_dir):
    return Path(cache_dir) / url.rstrip("/").rsplit("/", 1)[-1]


def _download(session, url, cache_dir):
    path = _cache_file(url, cache_dir)
    if path.exists() and path.stat().st_size > 0:
        return path
    resp = session.get(url)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(resp.content)
    return path


def find_tables_post(feed_xml):
    """First (newest) feed item whose title is a '... Situation-Tables (...)' release -> its link."""
    for item in re.findall(r"<item>(.*?)</item>", feed_xml, flags=re.S):
        title = html.unescape(re.search(r"<title>(.*?)</title>", item, re.S).group(1))
        if re.search(r"situation\s*[-–]\s*tables", title, re.I):
            link = html.unescape(re.search(r"<link>(.*?)</link>", item, re.S).group(1)).split("?")[0]
            return title, link
    return None


def newest_bfs_xlsx(page_html):
    """Pick the newest '<BSMonth>_<BSYear>...xlsx' link on the monthly-statistics page."""
    best = None
    for url in set(re.findall(r"https?://[^\"'\s<>]+?\.xlsx", page_html)):
        m = re.search(r"/([A-Za-z]+)[_\- ](\d{4})[^/]*\.xlsx$", url)
        if not m or m.group(1).lower() not in BS_MONTHS:
            continue
        ref = bs_month_to_reference(int(m.group(2)), m.group(1))
        if best is None or ref > best[0]:
            best = (ref, url)
    return best


def fetch_tables_workbook(session, cache_dir, offline=False, max_pages=5):
    cache_dir = Path(cache_dir)
    if offline:
        files = sorted(cache_dir.glob("tables/*.xlsx"), key=lambda p: p.stat().st_mtime)
        if not files:
            raise FileNotFoundError("no cached Tables workbook")
        return files[-1]
    for page in range(1, max_pages + 1):
        feed = _cached_text(session, CMES_FEED.format(page=page), cache_dir / f"cmes_feed_{page}.xml", offline)
        found = find_tables_post(feed)
        if found:
            title, link = found
            # The post link 302-redirects to the xlsx; file names are irregular, so follow it.
            xlsx_url = session.resolve(link)
            if not xlsx_url.lower().endswith((".xlsx", ".xls")):
                raise RuntimeError(f"Tables post {link} did not redirect to a spreadsheet: {xlsx_url}")
            print(f"Tables release: {title}\n  {xlsx_url}")
            return _download(session, xlsx_url, cache_dir / "tables")
    raise RuntimeError("no 'Tables' release found in the NRB feed")


def fetch_bfs_workbook(session, cache_dir, offline=False):
    cache_dir = Path(cache_dir)
    page = _cached_text(session, BFS_PAGE, cache_dir / "monthly_statistics.html", offline)
    best = newest_bfs_xlsx(page)
    if best is None:
        raise RuntimeError("no Banking and Financial Statistics xlsx link found")
    ref, url = best
    url = urljoin(BASE, url)
    print(f"BFS release: {ref.date()} reference month\n  {url}")
    if offline:
        return _cache_file(url, cache_dir / "bfs")
    return _download(session, url, cache_dir / "bfs")


def main(argv=None):
    import openpyxl

    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--cache-dir", default=str(DEFAULT_CACHE))
    parser.add_argument("--out", default=str(DEFAULT_OUT))
    parser.add_argument("--offline", action="store_true", help="Use cached files only")
    args = parser.parse_args(argv)

    session = None if args.offline else PoliteSession()
    tables_path = fetch_tables_workbook(session, args.cache_dir, args.offline)
    bfs_path = fetch_bfs_workbook(session, args.cache_dir, args.offline)
    tables_wb = openpyxl.load_workbook(tables_path, read_only=True, data_only=True)
    bfs_wb = openpyxl.load_workbook(bfs_path, read_only=True, data_only=True)
    df = build_macro(tables_wb, bfs_wb)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    df.round(4).to_csv(out, index=False, date_format="%Y-%m-%d")
    print(f"Wrote {len(df)} months ({df['month'].min().date()} .. {df['month'].max().date()}) to {out}")
    return df


if __name__ == "__main__":
    main()
